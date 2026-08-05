"""Importação da planilha 'CONTROLE DE POPs ELABORADOS E ATUALIZADOS'.

A planilha não é uma tabela simples: os POPs novos e atualizados ficam em
blocos com colunas agrupadas por ano (ex.: '2025 (60)' → par 'Nº POP | Nome').
Este parser localiza os blocos por marcadores de seção e lê os pares
(número, nome) de cada ano. A área de cada POP não existe por linha na
planilha (só há um resumo agregado), então é inferida por palavra-chave.

Upsert idempotente pela chave natural (numero|tipo|ano).
"""
import hashlib
import re
import unicodedata
from io import BytesIO

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .models import Pop

NPOP_SYN = {"n pop", "no pop", "nº pop", "num pop", "numero pop", "n. pop"}

# Classificador de área por palavra-chave (baseado no "Resumo por Área" da
# própria planilha). A primeira regra que casar vence.
AREA_REGRAS: list[tuple[str, str]] = [
    (r"qpcr|\bpcr\b|dna|sequenciamento|beads|buffer|leishm|cinomo|babesia|ehrlich|tricho|molecular|probe|sybr", "Biologia Molecular"),
    (r"fluxograma", "Fluxogramas Operacionais"),
    (r"conv[eê]nio|petlove|doutor pet|doglife|lifepet|medicina com carinho|petload", "Convênios e Atendimento"),
    (r"processo seletiv|contrata|integra|acompanhamento|gest[aã]o|recursos humanos|\brh\b|colaborad|localiza[cç][aã]o de documento", "Gestão Administrativa e RH"),
    (r"hematolog|imunolog|immulite|immunomat|advia|microscan|leuc[oó]citos|aglutina|rea[cç][aã]o cruzada", "Imunologia e Hematologia"),
    (r"bioqu[ií]mic|urin[aá]lise|urina|au480", "Bioquímica e Urinálise"),
    (r"centr[ií]fug|conting[eê]ncia|gerador|energia|descontamina|autoclave|transdutor|rss|res[ií]duo|infraestrutura|log[ií]stica|etiqueta", "Logística e Infraestrutura"),
    (r"equipamento|calibra|manuten|immulite|bionote|nanodrop|spectrum|aria|dry block|corador|advia|au480|microscan|adva|immunomat|eti-max|autoscan", "Equipamentos Laboratoriais"),
]


def _key(s) -> str:
    if s is None:
        return ""
    s = "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _norm(s) -> str:
    if s is None:
        return ""
    if isinstance(s, float) and s.is_integer():
        return str(int(s))
    return re.sub(r"\s+", " ", str(s)).strip()


def _ano(v) -> int | None:
    m = re.search(r"(20\d{2})", str(v) if v is not None else "")
    return int(m.group(1)) if m else None


def classificar_area(nome: str) -> str:
    n = _key(nome)
    for padrao, area in AREA_REGRAS:
        if re.search(padrao, n):
            return area
    return ""


def _chave(numero: str, tipo: str, ano) -> str:
    raw = f"{numero}|{tipo}|{ano or ''}"
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:32]


def _linha_marcador(rows, alvo: str, start: int = 0) -> int:
    for i in range(start, len(rows)):
        for c in rows[i]:
            if _key(c) == alvo:
                return i
    return -1


def _ler_bloco(rows, ini: int, fim: int, tipo: str) -> list[dict]:
    """Lê um bloco (novos/atualizados): acha a linha de cabeçalho com 'Nº POP',
    identifica os pares (col_numero, col_nome) e o ano de cada par (label acima),
    e coleta os itens das linhas de dados."""
    itens: list[dict] = []
    # Cabeçalho do bloco.
    h = -1
    for i in range(ini, fim):
        if any(_key(c) in NPOP_SYN for c in rows[i]):
            h = i
            break
    if h < 0:
        return itens
    header = rows[h]
    pares: list[tuple[int, int, int | None]] = []
    for c, cell in enumerate(header):
        if _key(cell) in NPOP_SYN:
            # Ano: procura um rótulo com ano nas linhas acima, na mesma coluna.
            ano = None
            for up in range(h - 1, max(ini - 1, h - 4), -1):
                ano = _ano(rows[up][c]) if c < len(rows[up]) else None
                if ano:
                    break
            pares.append((c, c + 1, ano))
    for r in range(h + 1, fim):
        row = rows[r]
        for col_num, col_nome, ano in pares:
            numero = _norm(row[col_num]) if col_num < len(row) else ""
            nome = _norm(row[col_nome]) if col_nome < len(row) else ""
            if not numero and not nome:
                continue
            if not numero:
                # sem número não há chave estável; ignora
                continue
            itens.append({
                "numero": numero, "nome": nome, "ano": ano, "tipo": tipo,
                "area": classificar_area(nome),
            })
    return itens


def importar_pops_xlsx(db: Session, conteudo: bytes) -> dict:
    """Lê o .xlsx de controle de POPs e faz upsert. Retorna {importados, total}."""
    wb = openpyxl.load_workbook(BytesIO(conteudo), data_only=True, read_only=True)
    # A aba de detalhe é a que tem os blocos 'POPs NOVOS'/'ATUALIZADOS'.
    ws = None
    for cand in wb.worksheets:
        vals = [_key(c) for row in cand.iter_rows(min_row=1, max_row=120, values_only=True) for c in row]
        if "pops novos" in vals or "pop novos" in vals:
            ws = cand
            break
    if ws is None:
        ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    i_novos = _linha_marcador(rows, "pops novos")
    # "POPs Atualizados" também aparece como cabeçalho do RESUMO GERAL no topo;
    # o título da seção é o que vem DEPOIS do bloco de novos.
    i_atual = _linha_marcador(rows, "pops atualizados", start=max(i_novos + 1, 0))
    # Fim dos blocos: linha do "Resumo por Área" (após atualizados).
    n = len(rows)
    i_area = -1
    for i in range(max(i_atual + 1, 0), n):
        if rows[i] and "resumo" in _key(rows[i][0]) and "area" in _key(rows[i][0]):
            i_area = i
            break
    fim_novos = i_atual if i_atual > i_novos else n
    fim_atual = i_area if i_area > i_atual else n

    itens: list[dict] = []
    if i_novos >= 0:
        itens += _ler_bloco(rows, i_novos + 1, fim_novos, "novo")
    if i_atual >= 0:
        itens += _ler_bloco(rows, i_atual + 1, fim_atual, "atualizado")

    importados = 0
    staged: dict[str, Pop] = {}
    for it in itens:
        chave = _chave(it["numero"], it["tipo"], it["ano"])
        existente = staged.get(chave) or db.scalar(select(Pop).where(Pop.chave == chave))
        if existente:
            existente.nome = it["nome"] or existente.nome
            existente.ano = it["ano"]
            # Preserva área já ajustada manualmente; só preenche se vazia.
            if not existente.area and it["area"]:
                existente.area = it["area"]
        else:
            novo = Pop(chave=chave, **it)
            db.add(novo)
            staged[chave] = novo
            importados += 1

    db.commit()
    total = db.scalar(select(func.count()).select_from(Pop)) or 0
    return {"importados": importados, "total": int(total)}
