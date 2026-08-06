"""Geração de relatórios (Excel e PDF) dos módulos POP e Chamados.

Excel via openpyxl; PDF via reportlab (Python puro, sem dependências nativas).
Formatação com a identidade Alchemypet (marinho/teal/dourado).
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .models import Chamado, Pop

# ===== Paleta =====
MARINHO = "1E3A5F"
TEAL = "0FA3A3"
GOLD = "C9A24B"
CINZA = "5A6B7A"
LINHA = "DDE5EC"
ZEBRA = "F4F7FA"
PIE_CORES = [
    "#3d6fb4", "#2e9e6b", "#d98a2b", "#c0532f", "#7a5aa8", "#0fa3a3",
    "#c9a24b", "#b4488f", "#4b8f3d", "#2b6f8f", "#a9600c", "#5a6b7a", "#8f8f2b",
]
PIE_DEMAIS = "#9aa7b4"

_ROTULO_COMPLEX = {"alta": "Alta", "media": "Média", "baixa": "Baixa"}
_ORDEM_COMPLEX = {"alta": 0, "media": 1, "baixa": 2}


def _hoje() -> str:
    return datetime.now().strftime("%d/%m/%Y %H:%M")


# ===================================================================
#  Coleta de dados
# ===================================================================
def _dados_pop(db: Session) -> dict:
    pops = list(db.scalars(select(Pop).order_by(Pop.ano.desc(), Pop.numero)))
    total = len(pops)
    novos = sum(1 for p in pops if p.tipo == "novo")
    atualizados = sum(1 for p in pops if p.tipo == "atualizado")
    anos = sorted({p.ano for p in pops if p.ano is not None})
    periodo = ""
    if anos:
        periodo = f"{anos[0]}" if anos[0] == anos[-1] else f"{anos[0]}–{anos[-1]}"
    por_ano = []
    for a in anos:
        nv = sum(1 for p in pops if p.ano == a and p.tipo == "novo")
        at = sum(1 for p in pops if p.ano == a and p.tipo == "atualizado")
        por_ano.append((a, nv, at, nv + at))
    areas: dict[str, int] = {}
    for p in pops:
        k = p.area or "Sem classificação"
        areas[k] = areas.get(k, 0) + 1
    por_area = sorted(areas.items(), key=lambda kv: -kv[1])
    return {
        "total": total, "novos": novos, "atualizados": atualizados,
        "periodo": periodo, "por_ano": por_ano, "por_area": por_area, "lista": pops,
    }


def _dados_chamados(db: Session) -> dict:
    chs = list(db.scalars(select(Chamado).order_by(Chamado.data.desc().nullslast(), Chamado.created_at.desc())))
    total = len(chs)
    resolvidos = sum(1 for c in chs if c.status == "resolvido")
    abertos = total - resolvidos
    taxa = round(100.0 * resolvidos / total, 1) if total else 0.0
    motivos: dict[str, int] = {}
    for c in chs:
        k = c.motivo or "Sem motivo"
        motivos[k] = motivos.get(k, 0) + 1
    por_motivo = sorted(motivos.items(), key=lambda kv: -kv[1])
    comp: dict[str, int] = {}
    for c in chs:
        comp[c.complexidade] = comp.get(c.complexidade, 0) + 1
    por_complex = [
        (_ROTULO_COMPLEX.get(k, k or "—"), v)
        for k, v in sorted(comp.items(), key=lambda kv: _ORDEM_COMPLEX.get(kv[0], 9))
    ]
    return {
        "total": total, "abertos": abertos, "resolvidos": resolvidos, "taxa": taxa,
        "por_motivo": por_motivo, "por_complexidade": por_complex, "lista": chs,
    }


# ===================================================================
#  Excel (openpyxl)
# ===================================================================
def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


_BORDA = Border(*[Side(style="thin", color=LINHA)] * 4)


def _titulo_aba(ws, titulo: str, ncols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, "ALCHEMYPET")
    c.font = Font(bold=True, size=11, color="FFFFFF")
    c.fill = _fill(MARINHO)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 22
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    t = ws.cell(2, 1, titulo)
    t.font = Font(bold=True, size=14, color=MARINHO)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    s = ws.cell(3, 1, f"Gerado em {_hoje()}")
    s.font = Font(size=9, color=CINZA)


def _cabecalho_tabela(ws, row: int, headers: list[str]):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row, j, h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = _fill(MARINHO)
        c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        c.border = _BORDA
    ws.row_dimensions[row].height = 18


def _linha_tabela(ws, row: int, valores: list, zebra: bool):
    for j, v in enumerate(valores, start=1):
        c = ws.cell(row, j, v)
        c.border = _BORDA
        c.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=False)
        if zebra:
            c.fill = _fill(ZEBRA)


def _kpis_excel(ws, row: int, kpis: list[tuple[str, str]]):
    for i, (label, valor) in enumerate(kpis):
        col = 1 + i * 2
        lab = ws.cell(row, col, label.upper())
        lab.font = Font(size=8, bold=True, color=CINZA)
        val = ws.cell(row + 1, col, valor)
        val.font = Font(size=16, bold=True, color=MARINHO)
    return row + 2


def _autolargura(ws, larguras: dict[int, int]):
    for col, w in larguras.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _barra_excel(ws, titulo, data_ref, cats_ref, anchor, tipo="col", alt=7.5, larg=13):
    ch = BarChart()
    ch.type = tipo
    ch.grouping = "clustered"
    ch.title = titulo
    ch.height = alt
    ch.width = larg
    ch.gapWidth = 60
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)


def _pizza_excel(ws, titulo, data_ref, cats_ref, anchor, alt=8, larg=13):
    p = PieChart()
    p.title = titulo
    p.height = alt
    p.width = larg
    p.add_data(data_ref, titles_from_data=True)
    p.set_categories(cats_ref)
    p.dataLabels = DataLabelList()
    p.dataLabels.showPercent = True
    ws.add_chart(p, anchor)


def pops_xlsx(db: Session) -> bytes:
    d = _dados_pop(db)
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumo"
    _titulo_aba(ws, "Controle de POPs — Resumo gerencial", 6)
    r = _kpis_excel(ws, 5, [
        ("Total de POPs", str(d["total"])),
        ("Novos", str(d["novos"])),
        ("Atualizados", str(d["atualizados"])),
        ("Período", d["periodo"] or "—"),
    ])
    r += 2
    ws.cell(r, 1, "POPs por ano").font = Font(bold=True, size=12, color=MARINHO)
    r += 1
    r_ano = r
    _cabecalho_tabela(ws, r, ["Ano", "Novos", "Atualizados", "Total"])
    n_ano = len(d["por_ano"])
    for i, (ano, nv, at, tot) in enumerate(d["por_ano"]):
        _linha_tabela(ws, r + 1 + i, [ano, nv, at, tot], i % 2 == 1)
    r = r + 1 + n_ano + 2
    ws.cell(r, 1, "POPs por área").font = Font(bold=True, size=12, color=MARINHO)
    r += 1
    r_area = r
    _cabecalho_tabela(ws, r, ["Área", "Quantidade"])
    n_area = len(d["por_area"])
    for i, (area, q) in enumerate(d["por_area"]):
        _linha_tabela(ws, r + 1 + i, [area, q], i % 2 == 1)
    _autolargura(ws, {1: 30, 2: 14, 3: 14, 4: 10})

    # Gráficos (à direita das tabelas).
    if n_ano:
        _barra_excel(
            ws, "POPs por ano (novos × atualizados)",
            Reference(ws, min_col=2, max_col=3, min_row=r_ano, max_row=r_ano + n_ano),
            Reference(ws, min_col=1, min_row=r_ano + 1, max_row=r_ano + n_ano),
            "F5",
        )
    if n_area:
        _barra_excel(
            ws, "POPs por área",
            Reference(ws, min_col=2, min_row=r_area, max_row=r_area + n_area),
            Reference(ws, min_col=1, min_row=r_area + 1, max_row=r_area + n_area),
            "F22", tipo="bar", alt=9,
        )

    ws2 = wb.create_sheet("POPs")
    _titulo_aba(ws2, "Controle de POPs — Lista completa", 5)
    _cabecalho_tabela(ws2, 5, ["Nº POP", "Nome", "Ano", "Tipo", "Área"])
    for i, p in enumerate(d["lista"]):
        _linha_tabela(ws2, 6 + i, [
            p.numero, p.nome, p.ano, "Novo" if p.tipo == "novo" else "Atualizado", p.area,
        ], i % 2 == 1)
    ws2.freeze_panes = "A6"
    _autolargura(ws2, {1: 10, 2: 58, 3: 8, 4: 14, 5: 28})

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def chamados_xlsx(db: Session) -> bytes:
    d = _dados_chamados(db)
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumo"
    _titulo_aba(ws, "Chamados por e-mail — Resumo gerencial", 6)
    r = _kpis_excel(ws, 5, [
        ("Total", str(d["total"])),
        ("Abertos", str(d["abertos"])),
        ("Resolvidos", str(d["resolvidos"])),
        ("Taxa de resolução", f"{d['taxa']}%"),
    ])
    r += 2
    ws.cell(r, 1, "Solicitações por motivo").font = Font(bold=True, size=12, color=MARINHO)
    r += 1
    r_mot = r
    _cabecalho_tabela(ws, r, ["Motivo", "Quantidade", "%"])
    soma = d["total"] or 1
    n_mot = len(d["por_motivo"])
    for i, (m, q) in enumerate(d["por_motivo"]):
        _linha_tabela(ws, r + 1 + i, [m, q, f"{round(100 * q / soma, 1)}%"], i % 2 == 1)
    r = r + 1 + n_mot + 2
    ws.cell(r, 1, "Por complexidade").font = Font(bold=True, size=12, color=MARINHO)
    r += 1
    r_cx = r
    _cabecalho_tabela(ws, r, ["Complexidade", "Quantidade", "%"])
    n_cx = len(d["por_complexidade"])
    for i, (c, q) in enumerate(d["por_complexidade"]):
        _linha_tabela(ws, r + 1 + i, [c, q, f"{round(100 * q / soma, 1)}%"], i % 2 == 1)
    _autolargura(ws, {1: 38, 2: 14, 3: 8})

    if n_mot:
        _pizza_excel(
            ws, "Distribuição por motivo",
            Reference(ws, min_col=2, min_row=r_mot, max_row=r_mot + n_mot),
            Reference(ws, min_col=1, min_row=r_mot + 1, max_row=r_mot + n_mot),
            "E5", alt=9, larg=15,
        )
    if n_cx:
        _barra_excel(
            ws, "Por complexidade",
            Reference(ws, min_col=2, min_row=r_cx, max_row=r_cx + n_cx),
            Reference(ws, min_col=1, min_row=r_cx + 1, max_row=r_cx + n_cx),
            "E24",
        )

    ws2 = wb.create_sheet("Chamados")
    _titulo_aba(ws2, "Chamados por e-mail — Lista completa", 5)
    _cabecalho_tabela(ws2, 5, ["Assunto", "Complexidade", "Motivo", "Status", "Link Gmail"])
    for i, c in enumerate(d["lista"]):
        _linha_tabela(ws2, 6 + i, [
            c.assunto, _ROTULO_COMPLEX.get(c.complexidade, c.complexidade),
            c.motivo, "Resolvido" if c.status == "resolvido" else "Aberto", c.link_gmail,
        ], i % 2 == 1)
    ws2.freeze_panes = "A6"
    _autolargura(ws2, {1: 40, 2: 14, 3: 34, 4: 12, 5: 40})

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ===================================================================
#  PDF (reportlab)
# ===================================================================
def _estilos():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("H1x", parent=ss["Heading1"], textColor=colors.HexColor("#" + MARINHO), fontSize=16, spaceAfter=2))
    ss.add(ParagraphStyle("Subx", parent=ss["Normal"], textColor=colors.HexColor("#" + CINZA), fontSize=8.5, spaceAfter=8))
    ss.add(ParagraphStyle("Secx", parent=ss["Heading2"], textColor=colors.HexColor("#" + MARINHO), fontSize=12, spaceBefore=8, spaceAfter=4))
    ss.add(ParagraphStyle("Cellx", parent=ss["Normal"], fontSize=8.5, leading=11))
    return ss


def _faixa_marca(largura):
    t = Table([["ALCHEMYPET"]], colWidths=[largura])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + MARINHO)),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return t


def _kpi_cards(kpis, largura):
    row = [f"{v}\n{lab}" for lab, v in kpis]
    # Duas linhas por célula: valor grande + label pequeno via Paragraph
    ss = _estilos()
    cells = []
    for lab, v in kpis:
        p = Paragraph(
            f'<font size=17 color="#{MARINHO}"><b>{v}</b></font><br/>'
            f'<font size=7.5 color="#{CINZA}">{lab.upper()}</font>',
            ss["Cellx"],
        )
        cells.append(p)
    n = len(cells)
    t = Table([cells], colWidths=[largura / n] * n)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F7FA")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#" + LINHA)),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return t


def _tabela(headers, linhas, larguras, ss, aligns=None):
    dados = [[Paragraph(f"<b>{h}</b>", ParagraphStyle("h", parent=ss["Cellx"], textColor=colors.white)) for h in headers]]
    for ln in linhas:
        dados.append([Paragraph(str(v) if v is not None else "", ss["Cellx"]) for v in ln])
    t = Table(dados, colWidths=larguras, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + MARINHO)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + ZEBRA)]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#" + LINHA)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for col, al in (aligns or {}).items():
        estilo.append(("ALIGN", (col, 0), (col, -1), al))
    t.setStyle(TableStyle(estilo))
    return t


def _grafico_pizza(itens, max_fatias=6):
    """Donut de participação por motivo (top N + Demais)."""
    dados = sorted(itens, key=lambda kv: -kv[1])
    if len(dados) > max_fatias:
        cabeca = dados[:max_fatias]
        resto = sum(v for _, v in dados[max_fatias:])
        dados = cabeca + ([("Demais", resto)] if resto else [])
    soma = sum(v for _, v in dados) or 1

    # Drawing compacto para caber na coluna (evita sobrepor a legenda).
    d = Drawing(150, 150)
    pie = Pie()
    pie.x = 8
    pie.y = 6
    pie.width = 134
    pie.height = 134
    pie.data = [v for _, v in dados]
    pie.innerRadiusFraction = 0.55
    pie.slices.strokeWidth = 0.5
    pie.slices.strokeColor = colors.white
    for i, (nome, v) in enumerate(dados):
        cor = PIE_DEMAIS if nome == "Demais" else PIE_CORES[i % len(PIE_CORES)]
        pie.slices[i].fillColor = colors.HexColor(cor)
    d.add(pie)
    return d, dados, soma


def _barras_pdf_ano(por_ano, largura):
    """Barras agrupadas (novos × atualizados) por ano para o PDF de POPs."""
    d = Drawing(largura, 150)
    bc = VerticalBarChart()
    bc.x = 34
    bc.y = 22
    bc.width = largura - 70
    bc.height = 110
    bc.data = [
        [nv for _, nv, _, _ in por_ano],
        [at for _, _, at, _ in por_ano],
    ]
    bc.categoryAxis.categoryNames = [str(a) for a, _, _, _ in por_ano]
    bc.valueAxis.valueMin = 0
    bc.groupSpacing = 18
    bc.barSpacing = 2
    bc.bars[0].fillColor = colors.HexColor("#2e9e6b")  # novos
    bc.bars[1].fillColor = colors.HexColor("#d98a2b")  # atualizados
    bc.bars.strokeColor = None
    d.add(bc)
    # Legenda manual (dois quadradinhos + rótulos).
    from reportlab.graphics.shapes import Rect, String as _Str
    d.add(Rect(largura - 150, 132, 9, 9, fillColor=colors.HexColor("#2e9e6b"), strokeColor=None))
    d.add(_Str(largura - 137, 133, "Novos", fontSize=8, fillColor=colors.HexColor("#" + CINZA)))
    d.add(Rect(largura - 95, 132, 9, 9, fillColor=colors.HexColor("#d98a2b"), strokeColor=None))
    d.add(_Str(largura - 82, 133, "Atualizados", fontSize=8, fillColor=colors.HexColor("#" + CINZA)))
    return d


def _pdf_doc(titulo: str, elementos: list, paisagem=False) -> bytes:
    out = BytesIO()
    tam = landscape(A4) if paisagem else A4
    doc = SimpleDocTemplate(
        out, pagesize=tam, topMargin=14 * mm, bottomMargin=14 * mm,
        leftMargin=14 * mm, rightMargin=14 * mm, title=titulo,
    )
    doc.build(elementos)
    return out.getvalue()


def pops_pdf(db: Session) -> bytes:
    d = _dados_pop(db)
    ss = _estilos()
    W = A4[0] - 28 * mm
    el = [
        _faixa_marca(W), Spacer(1, 8),
        Paragraph("Controle de POPs", ss["H1x"]),
        Paragraph(f"Relatório gerencial · gerado em {_hoje()}", ss["Subx"]),
        _kpi_cards([
            ("Total de POPs", d["total"]), ("Novos", d["novos"]),
            ("Atualizados", d["atualizados"]), ("Período", d["periodo"] or "—"),
        ], W),
        Spacer(1, 10),
        Paragraph("POPs por ano", ss["Secx"]),
        _tabela(["Ano", "Novos", "Atualizados", "Total"],
                [[a, nv, at, tot] for a, nv, at, tot in d["por_ano"]],
                [W * 0.3, W * 0.23, W * 0.24, W * 0.23], ss),
        Spacer(1, 4),
        _barras_pdf_ano(d["por_ano"], W) if d["por_ano"] else Spacer(1, 0),
        Spacer(1, 8),
        Paragraph("POPs por área", ss["Secx"]),
        _tabela(["Área", "Quantidade"],
                [[a, q] for a, q in d["por_area"]],
                [W * 0.7, W * 0.3], ss),
        Spacer(1, 8),
        Paragraph("Lista de POPs", ss["Secx"]),
        _tabela(["Nº", "Nome", "Ano", "Tipo", "Área"],
                [[p.numero, p.nome, p.ano, "Novo" if p.tipo == "novo" else "Atualizado", p.area] for p in d["lista"]],
                [W * 0.07, W * 0.44, W * 0.09, W * 0.15, W * 0.25], ss),
    ]
    return _pdf_doc("Controle de POPs", el)


def chamados_pdf(db: Session) -> bytes:
    d = _dados_chamados(db)
    ss = _estilos()
    W = A4[0] - 28 * mm
    desenho, fatias, soma = _grafico_pizza(d["por_motivo"])
    # Legenda ao lado da pizza (tabela).
    leg_linhas = []
    for i, (nome, v) in enumerate(fatias):
        pct = f"{round(100 * v / soma, 1)}%"
        leg_linhas.append([nome, v, pct])
    col_pizza = 165
    col_leg = W - col_pizza
    leg = _tabela(["Motivo", "Qtd", "%"], leg_linhas, [col_leg - 108, 60, 44], ss)
    pizza_bloco = Table([[desenho, leg]], colWidths=[col_pizza, col_leg])
    pizza_bloco.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, 0), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 0),
    ]))

    el = [
        _faixa_marca(W), Spacer(1, 8),
        Paragraph("Chamados por e-mail", ss["H1x"]),
        Paragraph(f"Relatório gerencial · gerado em {_hoje()}", ss["Subx"]),
        _kpi_cards([
            ("Total", d["total"]), ("Abertos", d["abertos"]),
            ("Resolvidos", d["resolvidos"]), ("Taxa de resolução", f"{d['taxa']}%"),
        ], W),
        Spacer(1, 10),
        Paragraph("Distribuição das solicitações por motivo", ss["Secx"]),
        pizza_bloco,
        Spacer(1, 8),
        Paragraph("Por complexidade", ss["Secx"]),
        _tabela(["Complexidade", "Quantidade", "%"],
                [[c, q, f"{round(100 * q / (d['total'] or 1), 1)}%"] for c, q in d["por_complexidade"]],
                [W * 0.4, W * 0.3, W * 0.3], ss),
        Spacer(1, 8),
        Paragraph("Chamados por motivo (detalhamento)", ss["Secx"]),
        _tabela(["Motivo", "Quantidade", "%"],
                [[m, q, f"{round(100 * q / (d['total'] or 1), 1)}%"] for m, q in d["por_motivo"]],
                [W * 0.5, W * 0.25, W * 0.25], ss),
    ]
    return _pdf_doc("Chamados por e-mail", el)
