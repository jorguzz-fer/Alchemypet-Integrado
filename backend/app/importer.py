"""Importação da planilha de pendências (todas as abas mensais).

Detecta colunas pelo NOME do cabeçalho (não pela posição), resistente às
variações entre abas ao longo dos anos. Mesma lógica validada no protótipo
client-side. Upsert idempotente pela chave natural (aba|guia|paciente|info).
"""
import hashlib
import re
import unicodedata
from datetime import date, datetime

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .models import Clinica, Pendencia

MES_MAP = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}

FIELD_SYN = {
    "guia": ["guia"],
    "paciente": ["paciente", "pacinte"],
    "cod_clinica": ["codigo da clinica", "cod da clinica", "cod. clinica", "cod clinica"],
    "clinica": ["clinica"],
    "info": ["informacao necessaria", "acao/comentario atendimento", "acao/comentario-atendimento", "acao/comentario"],
    "responsavel": ["responsavel"],
    "resposta": ["resposta do cliente"],
    "data_dev": ["data da devolutiva", "data da devulativa", "data devolutiva"],
    "colaborador": ["colaborador do atendimento", "colaborador"],
    "confirmacao": ["confirmacao para a clinica", "confitmacao p/ a clinica", "confirmacao p/ a clinica", "confirmacao"],
    "triagem": ["triagem"],
    "status": ["status"],
    "data_pedido": ["data do pedido da inf.", "data do pedido inf.", "data do pedido da inf", "data da solicitacao", "data de entrada"],
    "data_guia": ["data da guia"],
    "data_gen": ["data"],
}


def _norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def _key(s: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _fit(s: str, n: int) -> str:
    """Trunca ao tamanho máximo da coluna (evita StringDataRightTruncation)."""
    return s[:n] if len(s) > n else s


def _period(sheet_name: str):
    n = _key(sheet_name)
    m = re.search(r"([a-z]+)\s*0?(\d{4})", n)
    if m and MES_MAP.get(m.group(1)):
        return int(m.group(2)), MES_MAP[m.group(1)]
    m = re.search(r"([a-z]+?)0?(\d{4})", n)
    if m and MES_MAP.get(m.group(1)):
        return int(m.group(2)), MES_MAP[m.group(1)]
    return None, None


def _to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return None
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _build_colmap(header_row):
    colmap = {}
    for idx, cell in enumerate(header_row):
        h = _key(cell if cell is not None else "")
        if not h:
            continue
        for field, syns in FIELD_SYN.items():
            if field in colmap:
                continue
            if h in syns:
                colmap[field] = idx
                break
    return colmap


def _derive_status(confirmacao: str, resposta: str, data_dev) -> str:
    conf = _key(confirmacao)
    if conf == "ok" or conf.startswith("ok ") or conf == "okk":
        return "concluido"
    if resposta or data_dev:
        return "tratativa"
    return "pendente"


def _chave(aba, guia, paciente, info) -> str:
    raw = "|".join([aba, guia, paciente, info])
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:32]


def importar_xlsx(db: Session, conteudo: bytes) -> dict:
    """Lê o .xlsx, normaliza e faz upsert. Retorna {importados, abas, total}."""
    from io import BytesIO

    wb = openpyxl.load_workbook(BytesIO(conteudo), data_only=True, read_only=True)
    abas = 0
    importados = 0
    clinicas_vistas: dict[str, str] = {}  # nome -> codigo
    # Novas pendências deste lote ainda não commitadas (autoflush off):
    # rastreamos por chave para deduplicar dentro do próprio arquivo.
    staged: dict[str, Pendencia] = {}

    for sheet_name in wb.sheetnames:
        if re.match(r"^p[áa]gina", sheet_name, re.IGNORECASE):
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # localizar cabeçalho (contém "guia") nas primeiras 6 linhas
        h_idx = -1
        for i in range(min(6, len(rows))):
            if any(_key(c) == "guia" for c in rows[i] if c is not None):
                h_idx = i
                break
        if h_idx < 0:
            continue
        colmap = _build_colmap(rows[h_idx])
        if "guia" not in colmap and "paciente" not in colmap:
            continue

        ano, mes = _period(sheet_name)

        def get(row, field):
            i = colmap.get(field)
            return row[i] if i is not None and i < len(row) else None

        abas += 1
        for r in range(h_idx + 1, len(rows)):
            row = rows[r]
            guia = _norm(get(row, "guia"))
            paciente = _norm(get(row, "paciente"))
            if not guia and not paciente:
                continue

            info = _norm(get(row, "info"))
            resposta = _norm(get(row, "resposta"))
            confirmacao = _norm(get(row, "confirmacao"))
            data_pedido = _to_date(get(row, "data_pedido")) or _to_date(get(row, "data_guia")) or _to_date(get(row, "data_gen"))
            data_dev = _to_date(get(row, "data_dev"))
            clinica = _norm(get(row, "clinica"))
            cod_clinica = _norm(get(row, "cod_clinica"))
            status = _derive_status(confirmacao, resposta, data_dev)
            chave = _chave(sheet_name, guia, paciente, info)

            if clinica and clinica not in clinicas_vistas:
                clinicas_vistas[clinica] = cod_clinica

            # Trunca aos limites das colunas varchar (o Postgres é rígido).
            # Necessário porque abas com colunas desalinhadas na origem podem
            # jogar um comentário longo num campo curto (ex.: confirmação).
            campos = dict(
                guia=_fit(guia, 40), paciente=_fit(paciente, 160),
                cod_clinica=_fit(cod_clinica, 40), clinica=_fit(clinica, 200),
                informacao_necessaria=info, resposta_cliente=resposta,
                responsavel=_fit(_norm(get(row, "responsavel")), 120),
                colaborador=_fit(_norm(get(row, "colaborador")), 120),
                confirmacao=_fit(confirmacao, 120), triagem=_fit(_norm(get(row, "triagem")), 120),
                status=status, ano=ano, mes=mes, data_pedido=data_pedido, data_devolutiva=data_dev,
                aba=_fit(sheet_name, 60),
            )

            existente = staged.get(chave) or db.scalar(select(Pendencia).where(Pendencia.chave == chave))
            if existente:
                # Preserva a gestão (controlada no painel); atualiza dados da planilha.
                # Duplicatas dentro do mesmo arquivo: o último registro prevalece.
                for k, v in campos.items():
                    setattr(existente, k, v)
            else:
                # Gestão inicial deriva do status (concluído -> resolvido).
                gestao = "resolvido" if status == "concluido" else "aberto"
                nova = Pendencia(chave=chave, gestao=gestao, **campos)
                db.add(nova)
                staged[chave] = nova
                importados += 1

    wb.close()

    # Upsert de master data de clínicas
    for nome, codigo in clinicas_vistas.items():
        nome = _fit(nome, 200)
        existe = db.scalar(select(Clinica).where(Clinica.nome == nome))
        if not existe:
            db.add(Clinica(nome=nome, codigo=_fit(codigo, 40)))

    db.commit()
    total = db.scalar(select(func.count()).select_from(Pendencia))
    return {"importados": importados, "abas": abas, "total": int(total or 0)}
